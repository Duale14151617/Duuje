"""Food stores, bulk entry: opening stock from a spreadsheet, and one Army order issued to many units at once.

Both go through the same rules and the same append-only ledger as the single-entry screens, and both are all-or-nothing: if any
row has a problem, nothing is saved and every problem is reported by row number.

Opening stock: a CSV or XLSX file, one row per batch. A store (unit + food category) may be given an opening balance only while it
has no stock movements at all, so a file cannot be imported twice and cannot overwrite a running ledger. Each store gets ONE
"opening" ledger entry (the ledger allows only one per store) with the total, and one batch per row beneath it carrying its own
expiry date, supplier and batch reference, so expiry warnings and earliest-expiry-first issuing work from day one.

Bulk transfer: one source store, one Army order, many destination units. Each destination becomes an ordinary allocation that is
"on the way" until that unit's signed delivery note is recorded, exactly as for a single allocation.
"""
import csv
import io
import re
from collections import OrderedDict
from datetime import date, datetime

from django.db import IntegrityError, transaction
from django.db.models import F

from apps.accounts import scope
from apps.audit.services import record
from apps.core.errors import field_error
from apps.organization.models import OrganizationalUnit

from . import food, services as svc
from .models import FoodBatch, InventoryCategory, StockLine, StockTransaction, Supplier

MAX_BYTES = 1024 * 1024
MAX_ROWS = 1000
MAX_LINES = 100
MAX_QTY = 100_000_000
COLUMNS = ("unit_code", "category_code", "quantity", "uom", "expires_on", "received_on", "batch_ref", "supplier",
           "delivery_note_ref", "reorder_level")
REQUIRED = {"unit_code", "category_code", "quantity"}


# ------------------------------------------------------------------ reading the file
def _header(value):
    return re.sub(r"[\s\-]+", "_", str(value or "").strip().lower())


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):  # a real date typed into Excel arrives as a datetime
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():  # Excel stores 100 as 100.0
        return str(int(value))
    return str(value).strip()[:200]


def _read_csv(upload):
    try:
        text = upload.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        raise field_error("file", "file_encoding_invalid", "Save the file as UTF-8.")
    first = text.split("\n", 1)[0]
    delimiter = ";" if first.count(";") > first.count(",") else ","  # some regional Excel settings save with semicolons
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    table = []
    try:
        for cells in reader:
            if any(_cell(c) for c in cells):
                table.append((reader.line_num, cells))
            if len(table) > MAX_ROWS + 1:
                raise field_error("file", "file_too_many_rows", "Too many rows.")
    except csv.Error:
        raise field_error("file", "file_columns_invalid", "The columns do not match the template.")
    return table


def _read_xlsx(upload):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise field_error("file", "xlsx_unavailable", "Excel files are not enabled on this server. Save the sheet as CSV.")
    table, too_many = [], False
    try:
        wb = load_workbook(io.BytesIO(upload.read()), read_only=True, data_only=True)  # data_only: values, never formulas
        try:
            for number, cells in enumerate(wb.worksheets[0].iter_rows(values_only=True), start=1):
                if any(_cell(c) for c in cells):
                    table.append((number, list(cells)))
                if len(table) > MAX_ROWS + 1:
                    too_many = True
                    break
        finally:
            wb.close()
    except Exception:  # a damaged or unsupported workbook can fail in many ways; none of them is the user's data
        raise field_error("file", "food_file_type_invalid", "Upload a .csv or .xlsx file.")
    if too_many:
        raise field_error("file", "file_too_many_rows", "Too many rows.")
    return table


def parse_upload(upload):
    """The file as a list of row dicts (with the spreadsheet row number under "row"). Raises a field error for an unsafe or
    malformed file."""
    if upload.size > MAX_BYTES:
        raise field_error("file", "file_too_large", "The file is too large.")
    name = (upload.name or "").lower()
    if name.endswith(".csv"):
        table = _read_csv(upload)
    elif name.endswith(".xlsx"):
        table = _read_xlsx(upload)
    else:
        raise field_error("file", "food_file_type_invalid", "Upload a .csv or .xlsx file.")
    if not table:
        raise field_error("file", "file_empty", "The file has no rows.")
    names = [_header(c) for c in table[0][1]]
    known = [n for n in names if n]
    if not REQUIRED <= set(known) or not set(known) <= set(COLUMNS) or len(set(known)) != len(known):
        raise field_error("file", "file_columns_invalid", "The columns do not match the template.")
    rows = []
    for number, cells in table[1:]:
        row = {c: "" for c in COLUMNS}
        row["row"] = number
        for column, value in zip(names, cells):
            if column:
                row[column] = _cell(value)
        rows.append(row)
    if not rows:
        raise field_error("file", "file_empty", "The file has no rows.")
    return rows


# ------------------------------------------------------------------ checking the rows
def _whole(text, low, high):
    """A whole number. A comma is accepted only as a thousands separator (1,200), so a decimal comma (12,5) is refused
    instead of being read as 125."""
    digits = str(text).strip()
    if re.fullmatch(r"\d{1,3}(,\d{3})+", digits):
        digits = digits.replace(",", "")
    if not re.fullmatch(r"\d{1,12}", digits):
        return None
    n = int(digits)
    return n if low <= n <= high else None


def _day(text):
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def validate_opening(user, rows):
    """Check every row against the importing user's own scope and the stock ledger. Returns one result per row:
    {"row", "errors": [codes], ...resolved values...}. Writes nothing."""
    units = {u.code.upper(): u for u in scope.scoped_units(user, food.PERM)
             .filter(code__in={r["unit_code"].upper() for r in rows if r["unit_code"]})}
    cats = {c.code.lower(): c for c in InventoryCategory.objects.filter(
        code__in={r["category_code"].lower() for r in rows if r["category_code"]})}
    suppliers = {s.name.lower(): s for s in Supplier.objects.all()}
    can_restricted = {}

    pairs, ctx = set(), []
    for r in rows:
        unit, cat = units.get(r["unit_code"].upper()), cats.get(r["category_code"].lower())
        ctx.append((unit, cat))
        if unit is not None and cat is not None:
            pairs.add((unit.pk, cat.pk))
    lines = {(ln.org_unit_id, ln.category_id): ln for ln in StockLine.objects.filter(
        org_unit__in={p[0] for p in pairs}, category__in={p[1] for p in pairs}) if (ln.org_unit_id, ln.category_id) in pairs}
    moved = set(StockTransaction.objects.filter(stock_line__in=[ln.pk for ln in lines.values()])
                .values_list("stock_line_id", flat=True).distinct())

    results, seen, uoms, reorders, totals = [], set(), {}, {}, {}
    today = date.today()
    for r, (unit, cat) in zip(rows, ctx):
        errors = []
        if not r["unit_code"]:
            errors.append("unit_required")
        elif unit is None:
            errors.append("unit_unknown")  # unknown and out of scope look the same on purpose
        elif unit.level not in ("branch_command", "unit") or not unit.active:
            errors.append("food_unit_level_invalid")
        if not r["category_code"]:
            errors.append("category_required")
        elif cat is None:
            errors.append("food_category_unknown")
        elif not cat.code.startswith("food"):
            errors.append("not_food_category")
        elif cat.kind != "stock" or cat.approval_status != "approved" or not cat.active:
            errors.append("category_not_approved")
        elif cat.restricted and (unit is None or not can_restricted.setdefault(
                unit.pk, scope.can(user, "logistics.restricted", unit=unit))):
            errors.append("food_category_unknown")
        quantity = _whole(r["quantity"], 1, MAX_QTY)
        if quantity is None:
            errors.append("invalid_quantity")
        expires_on = received_on = None
        if r["expires_on"]:
            expires_on = _day(r["expires_on"])
            if expires_on is None:
                errors.append("date_invalid")
        if r["received_on"]:
            received_on = _day(r["received_on"])
            if received_on is None:
                errors.append("date_invalid")
            elif expires_on and expires_on < received_on:
                errors.append("expires_before_received")
        # Food that is already past its date can still be on the shelf, so an expiry in the past is fine when no delivery date is given.
        supplier = None
        if r["supplier"]:
            supplier = suppliers.get(r["supplier"].lower())
            if supplier is None:
                errors.append("supplier_unknown")
            elif not supplier.active:
                errors.append("supplier_inactive")
        if len(r["batch_ref"]) > 60:
            errors.append("batch_ref_too_long")
        if len(r["delivery_note_ref"]) > 60:
            errors.append("delivery_note_too_long")
        reorder = None
        if r["reorder_level"]:
            reorder = _whole(r["reorder_level"], 0, MAX_QTY)
            if reorder is None:
                errors.append("reorder_invalid")

        pair = (unit.pk, cat.pk) if unit is not None and cat is not None else None
        if pair is not None:
            line = lines.get(pair)
            if line is not None and line.pk in moved:
                errors.append("store_has_movements")
            if r["uom"]:
                if len(r["uom"]) > 16:
                    errors.append("uom_invalid")
                else:
                    known = uoms.get(pair) or (line.uom if line else None)
                    if known and known.lower() != r["uom"].lower():
                        errors.append("uom_mismatch")
                    else:
                        uoms.setdefault(pair, known or r["uom"])
            if reorder is not None:
                if reorders.setdefault(pair, reorder) != reorder:
                    errors.append("reorder_conflict")
            if quantity is not None:
                totals[pair] = totals.get(pair, 0) + quantity
                if totals[pair] > 2_000_000_000:
                    errors.append("invalid_quantity")
            if not errors:  # only a row that is otherwise fine can be a repeat, so one slip does not report a pile of duplicates
                key = (pair, r["batch_ref"].lower(), expires_on, supplier.pk if supplier else None, received_on,
                       r["delivery_note_ref"].lower())
                if key in seen:
                    errors.append("duplicate_batch_row")
                seen.add(key)
        results.append({"row": r["row"], "errors": errors, "pair": pair, "unit": unit, "cat": cat, "quantity": quantity,
                        "expires_on": expires_on, "received_on": received_on or today, "supplier": supplier,
                        "batch_ref": r["batch_ref"], "delivery_note_ref": r["delivery_note_ref"]})
    for res in results:
        res["uom"] = uoms.get(res["pair"]) or "kg"
        res["reorder"] = reorders.get(res["pair"])
    return results


def _stores(results):
    groups = OrderedDict()
    for res in results:
        groups.setdefault(res["pair"], []).append(res)
    return groups


def summarise(results):
    """One line per store for the preview, so the person can check the totals before confirming."""
    out = []
    for group in _stores(results).values():
        first = group[0]
        dates = [g["expires_on"] for g in group if g["expires_on"]]
        out.append({"unit_code": first["unit"].code, "unit_name": first["unit"].name_en, "category": first["cat"].name_en,
                    "uom": first["uom"], "quantity": sum(g["quantity"] for g in group), "batches": len(group),
                    "earliest_expiry": min(dates).isoformat() if dates else None})
    return out


# ------------------------------------------------------------------ saving
def commit_opening(request, results, authority_ref):
    """Create the opening balances. `results` must come from validate_opening() with no errors. All or nothing."""
    groups = _stores(results)
    for group in groups.values():
        svc.authorize(request, group[0]["unit"], group[0]["cat"], perm=food.PERM)
    user = request.user
    try:
        with transaction.atomic():
            for group in groups.values():
                first = group[0]
                unit, cat = first["unit"], first["cat"]
                line, _ = StockLine.objects.get_or_create(category=cat, org_unit=unit, defaults={
                    "uom": first["uom"], "reorder_level": first["reorder"] or 0})
                StockLine.objects.select_for_update().get(pk=line.pk)  # same lock the ledger uses
                if StockTransaction.objects.filter(stock_line=line).exists():  # someone moved stock since the check
                    raise field_error("file", "store_has_movements", "That store already has stock movements.")
                if first["reorder"] is not None and line.reorder_level == 0:
                    StockLine.objects.filter(pk=line.pk).update(reorder_level=first["reorder"], version=F("version") + 1)
                total = sum(g["quantity"] for g in group)
                txn = svc._apply(user, line.pk, "opening", total, "Opening balance (spreadsheet import)", authority_ref=authority_ref)
                for g in group:
                    FoodBatch.objects.create(
                        stock_line=line, supplier=g["supplier"], batch_ref=g["batch_ref"], quantity=g["quantity"],
                        remaining=g["quantity"], received_on=g["received_on"], expires_on=g["expires_on"], source="opening",
                        delivery_note_ref=g["delivery_note_ref"], authority_ref=authority_ref, txn=txn, created_by=user,
                        created_by_username=user.username)
                record("food.opening_balance", request=request, entity_type="stock_line", entity_id=line.pk,
                       metadata={"unit": str(unit.pk), "category": cat.code, "quantity": total, "batches": len(group),
                                 "authority_ref": authority_ref})
            record("food.opening_import", request=request, entity_type="food_batch",
                   metadata={"rows": len(results), "stores": len(groups), "authority_ref": authority_ref})
    except IntegrityError:  # the database's one-opening-per-store rule caught a concurrent import
        raise field_error("file", "import_conflict", "Another import changed these stores. Check the file again.")
    return {"rows": len(results), "stores": len(groups), "quantity": sum(r["quantity"] for r in results)}


# ------------------------------------------------------------------ bulk transfer
class BulkProblems(Exception):
    def __init__(self, rows):
        super().__init__("bulk_has_errors")
        self.rows = rows


def bulk_issue(request, d):
    """Issue one store's food to several units on one order. Every line is checked first and reported by line number; if all are
    fine, each becomes a normal allocation, all inside one transaction (so a failure part-way leaves nothing behind)."""
    line = StockLine.objects.select_related("category", "org_unit").filter(pk=str(d.get("stock_line") or "")).first() \
        if food._isid(d.get("stock_line")) else None
    if line is None:
        raise field_error("stock_line", "stock_line_invalid", "Choose a store.")
    food.food_category(line.category)
    svc.authorize(request, line.org_unit, line.category, perm=food.PERM)  # a refusal is audited and is not rolled back
    order_ref, officer, order_date = food._order(d)
    raw = d.get("lines")
    if not isinstance(raw, list) or not raw:
        raise field_error("lines", "lines_required", "Add at least one unit.")
    if len(raw) > MAX_LINES:
        raise field_error("lines", "too_many_lines", "Too many units in one order.")
    prefix = food._text(d.get("issue_note_prefix"), "issue_note_prefix", "", 50, required=False)

    used = set(StockTransaction.objects.filter(stock_line=line, kind="transfer_out").exclude(external_ref="")
               .values_list("external_ref", flat=True))
    wanted = [x.get("to_unit") for x in raw if isinstance(x, dict) and food._isid(x.get("to_unit"))]
    units = {str(u.pk): u for u in OrganizationalUnit.objects.filter(pk__in=wanted)}
    problems, plan, seen_units, seen_notes, running = [], [], set(), set(), 0
    for i, x in enumerate(raw, start=1):
        x = x if isinstance(x, dict) else {}
        errors = []
        unit = units.get(str(x.get("to_unit")))
        if unit is None or not scope.can(request.user, food.PERM, unit=unit):
            errors.append("org_unit_invalid")
        else:
            if unit.pk == line.org_unit_id:
                errors.append("no_change")
            elif unit.pk in seen_units:
                errors.append("duplicate_unit")
            if unit.level not in ("branch_command", "unit") or not unit.active:
                errors.append("food_unit_level_invalid")
            seen_units.add(unit.pk)
        quantity = _whole(x.get("quantity", ""), 1, MAX_QTY)
        if quantity is None:
            errors.append("invalid_quantity")
        elif running + quantity > line.quantity:
            errors.append("insufficient_stock")
        note = str(x.get("issue_note_ref") or "").strip() or (f"{prefix}-{i:02d}" if prefix else "")
        if not note:
            errors.append("issue_note_required")
        elif len(note) > 60:
            errors.append("issue_note_too_long")
        elif note in used or note in seen_notes:
            errors.append("duplicate_issue_note")
        seen_notes.add(note)
        if errors:
            problems.append({"row": i, "errors": errors})
        else:
            running += quantity
            plan.append((unit, quantity, note))
    if problems:
        raise BulkProblems(problems)

    allocations = []
    with transaction.atomic():
        for unit, quantity, note in plan:
            allocations.append(food.issue(request, {
                "stock_line": str(line.pk), "to_unit": str(unit.pk), "quantity": quantity, "order_ref": order_ref,
                "issuing_officer": officer, "order_date": order_date.isoformat(), "issue_note_ref": note}))
        record("food.bulk_issue", request=request, entity_type="stock_line", entity_id=line.pk,
               metadata={"lines": len(allocations), "quantity": running, "order_ref": order_ref,
                         "allocations": [str(a.pk) for a in allocations]})
    return allocations
